import datetime

from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect
from django.views import generic
from django.urls import reverse_lazy
from django.contrib import messages
from django.db import transaction
from django.utils import timezone

from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.contrib.messages.views import SuccessMessageMixin
from django.contrib.auth.decorators import login_required, permission_required

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import (
    Categoria, SubCategoria, Marca, UnidadMedida, Producto, TipoCambio, HistorialPrecioProducto,
    AjusteInventarioEnc, AjusteInventarioDet, MotivoAjusteInventario,
    StockSucursal, ajustar_stock_sucursal, PrecioSucursal,
    TransferenciaStockEnc, TransferenciaStockDet,
)
from .forms import CategoriaForm, SubCategoriaForm, MarcaForm, UnidadMedidaForm, ProductoForm, TipoCambioForm, \
    AjusteInventarioEncForm

from bases.alcance import requiere_alcance, AlcanceObjetoMixin
from bases.views import SinPrivilegios, obtener_sucursal_actual, es_casa_matriz

class CategoriaView(SinPrivilegios,generic.ListView):
    permission_required = "inv.view_categoria"
    model = Categoria
    template_name = "inv/categoria_list.html"
    context_object_name = "obj"
    
class CategoriaNew(SuccessMessageMixin,SinPrivilegios,generic.CreateView):
    permission_required = "inv.add_categoria"
    model = Categoria
    template_name = "inv/categoria_form.html"
    context_object_name = "obj"
    form_class = CategoriaForm
    success_url = reverse_lazy("inv:categoria_list")
    success_message="Categoria Creada Satisfactoriamente"

    def form_valid(self, form):
        form.instance.uc = self.request.user
        return super().form_valid(form)
    
class CategoriaEdit(SuccessMessageMixin,SinPrivilegios,generic.UpdateView):
    permission_required = "inv.change_categoria"
    model = Categoria
    template_name = "inv/categoria_form.html"
    context_object_name = "obj"
    form_class = CategoriaForm
    success_url = reverse_lazy("inv:categoria_list")
    login_url = "bases:login"
    success_message="Categoria Actualizada Satisfactoriamente"

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)
    
class CategoriaDel(SinPrivilegios, generic.DeleteView):
    permission_required = "inv.delete_categoria"
    model = Categoria
    template_name = "inv/catalogos_del.html"
    context_object_name = "obj"
    success_url = reverse_lazy("inv:categoria_list")

class SubCategoriaView(SinPrivilegios,generic.ListView):
    permission_required = "inv.view_subcategoria"
    model = SubCategoria
    template_name = "inv/subcategoria_list.html"
    context_object_name = "obj"

class SubCategoriaNew(SinPrivilegios, generic.CreateView):
    permission_required = "inv.add_subcategoria"
    model = SubCategoria
    template_name = "inv/subcategoria_form.html"
    context_object_name = "obj"
    form_class = SubCategoriaForm
    success_url = reverse_lazy("inv:subcategoria_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.uc = self.request.user
        return super().form_valid(form)
    
class SubCategoriaEdit(SinPrivilegios, generic.UpdateView):
    permission_required = "inv.change_subcategoria"
    model = SubCategoria
    template_name = "inv/subcategoria_form.html"
    context_object_name = "obj"
    form_class = SubCategoriaForm
    success_url = reverse_lazy("inv:subcategoria_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)
    
class SubCategoriaDel(SinPrivilegios, generic.DeleteView):
    permission_required = "inv.delete_subcategoria"
    model = SubCategoria
    template_name = "inv/subcatalogos_del.html"
    context_object_name = "obj"
    success_url = reverse_lazy("inv:subcategoria_list")

class MarcaView(SinPrivilegios,generic.ListView):
    permission_required = "inv.view_marca"
    model = Marca
    template_name = "inv/marca_list.html"
    context_object_name = "obj"

class MarcaNew(SinPrivilegios, generic.CreateView):
    permission_required = "inv.add_marca"
    model = Marca
    template_name = "inv/marca_form.html"
    context_object_name = "obj"
    form_class = MarcaForm
    success_url = reverse_lazy("inv:marca_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.uc = self.request.user
        return super().form_valid(form)
    
class MarcaEdit(SinPrivilegios, generic.UpdateView):
    permission_required = "inv.change_marca"
    model = Marca
    template_name = "inv/marca_form.html"
    context_object_name = "obj"
    form_class = MarcaForm
    success_url = reverse_lazy("inv:marca_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)

@login_required(login_url='/login/')
@permission_required('inv.change_marca',login_url='bases:sin_privilegios')
def marca_inactivar(request, id):
    marca = Marca.objects.filter(pk=id).first()
    contexto={}
    template_name="inv/catalogos_inactivo.html"

    if not marca:
        return redirect("inv:marca_list")

    if request.method=='GET':
        contexto={'obj':marca}
    
    if request.method=='POST':
        marca.estado=False
        marca.save()
        messages.success(request, 'Marca Inactivada.')
        return redirect("inv:marca_list")

    return render(request,template_name,contexto)

class UMView(SinPrivilegios, generic.ListView):
    permission_required = "inv.view_unidadmedida"
    model = UnidadMedida
    template_name = "inv/um_list.html"
    context_object_name = "obj"
    login_url = "bases:login"

class UMNew(SinPrivilegios, generic.CreateView):
    permission_required = "inv.add_unidadmedida"
    model = UnidadMedida
    template_name = "inv/um_form.html"
    context_object_name = "obj"
    form_class = UnidadMedidaForm
    success_url = reverse_lazy("inv:um_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.uc = self.request.user
        return super().form_valid(form)
    
class UMEdit(SinPrivilegios, generic.UpdateView):
    permission_required = "inv.change_unidadmedida"
    model = UnidadMedida
    template_name = "inv/um_form.html"
    context_object_name = "obj"
    form_class = UnidadMedidaForm
    success_url = reverse_lazy("inv:um_list")
    login_url = "bases:login"

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)

@login_required(login_url='/login/')
@permission_required('inv.change_unidadmedida',login_url='bases:sin_privilegios')
def um_inactivar(request, id):
    um = UnidadMedida.objects.filter(pk=id).first()
    contexto={}
    template_name="inv/catalogos_inactivo.html"

    if not um:
        return redirect("inv:um_list")

    if request.method=='GET':
        contexto={'obj':um}
    
    if request.method=='POST':
        um.estado=False
        um.save()
        return redirect("inv:um_list")

    return render(request,template_name,contexto)

class ProductoView(SinPrivilegios, generic.ListView):
    permission_required = "inv.view_producto"
    model = Producto
    template_name = "inv/producto_list.html"
    context_object_name = "obj"
    login_url = "bases:login"

    def get_context_data(self, **kwargs):
        # Existencia por sucursal (25/09/2026): con alcance limitado o un
        # sucursal elegida en el selector, se muestra el stock de ESAS
        # sucursales (StockSucursal), no el total de la empresa. Sin
        # restriccion ni seleccion, sigue mostrando el total.
        from bases.alcance import mapa_stock, contexto_filtro_sucursal
        context = super().get_context_data(**kwargs)
        mapa = mapa_stock(self.request)
        productos = list(context['obj'])
        if mapa is not None:
            for p in productos:
                p.existencia = mapa.get(p.id, 0)
        # Precio efectivo (25/09/2026): el local de la sucursal actual si
        # lo tiene, si no el base -- es lo que se factura en esa sucursal.
        _actual = obtener_sucursal_actual(self.request)
        if _actual is not None:
            _locales = dict(PrecioSucursal.objects.filter(sucursal=_actual).values_list('producto_id', 'precio'))
            for p in productos:
                if p.id in _locales:
                    p.precio = _locales[p.id]
        context['obj'] = productos
        context.update(contexto_filtro_sucursal(self.request))
        return context

def _modal_solo_central(request, mensaje):
    """Respuesta para pantallas que se abren como modal: un fragmento con
    el aviso (un redirect inyectaria toda la pagina dentro del modal)."""
    html = (
        '<div class="modal-dialog"><div class="modal-content"><div class="modal-body">'
        '<div class="alert alert-warning mb-3"><i class="fas fa-lock mr-2"></i>' + mensaje + '</div>'
        '<button type="button" class="btn btn-secondary" data-dismiss="modal">Cerrar</button>'
        '</div></div></div>'
    )
    return HttpResponse(html, status=200 if request.method == 'GET' else 403)


class ProductoNew(SuccessMessageMixin,SinPrivilegios, generic.CreateView):
    permission_required = "inv.add_producto"
    model = Producto
    template_name = "inv/producto_form.html"
    context_object_name = "obj"
    form_class = ProductoForm
    success_url = reverse_lazy("inv:producto_list")
    login_url = "bases:login"
    success_message="Producto Creado Satisfactoriamente"

    def dispatch(self, request, *args, **kwargs):
        # Solo la Central da de alta productos (catalogo unico de la
        # empresa, con su precio base) -- 25/09/2026, pedido de Carlos.
        if request.user.is_authenticated and not es_casa_matriz(request):
            return _modal_solo_central(
                request, 'Los productos nuevos (y sus precios) solo se dan de alta desde la Central.')
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.uc = self.request.user
        self.request.session['success_message'] = self.success_message

        redirect_url = reverse_lazy("inv:producto_list")

        response = JsonResponse({'redirect_url': redirect_url, 'success_message': self.success_message}, status=200)

        response['X-Redirect'] = redirect_url

        return super().form_valid(form)
    
    def get_context_data(self, **kwargs):
        context = super(ProductoNew, self).get_context_data(**kwargs)
        context["categorias"] = Categoria.objects.all().order_by('descripcion')
        context["subcategorias"] = SubCategoria.objects.all()
        return context
    
class ProductoEdit(SuccessMessageMixin,SinPrivilegios, generic.UpdateView):
    permission_required = "inv.change_producto"
    model = Producto
    template_name = "inv/producto_form.html"
    context_object_name = "obj"
    form_class = ProductoForm
    success_url = reverse_lazy("inv:producto_list")
    login_url = "bases:login"
    success_message="Producto Actualizado Satisfactoriamente"

    def get_form_kwargs(self):
        # Desde una sucursal los campos de precio quedan bloqueados (solo
        # la Central cambia precios) -- se hace cumplir en el servidor,
        # ver ProductoForm.
        kwargs = super().get_form_kwargs()
        kwargs['precios_editables'] = es_casa_matriz(self.request)
        return kwargs

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)
    
    def get_context_data(self, **kwargs):
        pk = self.kwargs.get('pk')
        
        context = super(ProductoEdit, self).get_context_data(**kwargs)
        context["categorias"] = Categoria.objects.all().order_by('descripcion')
        context["subcategorias"] = SubCategoria.objects.all()
        context["obj"] = Producto.objects.filter(pk=pk).first()

        return context

@login_required(login_url='/login/')
@permission_required('inv.change_producto',login_url='bases:sin_privilegios')    
def producto_inactivar(request, id):
    prod = Producto.objects.filter(pk=id).first()
    contexto={}
    template_name="inv/catalogos_inactivo.html"

    if not prod:
        return redirect("inv:producto_list")

    if request.method=='GET':
        contexto={'obj':prod}
    
    if request.method=='POST':
        prod.estado=False
        prod.save()
        return redirect("inv:producto_list")

    return render(request,template_name,contexto)


# =====================================================================
# Homologacion de Productos ante el SIN (RND 102500000018)
# =====================================================================

@login_required(login_url='/login/')
@permission_required('inv.change_producto', login_url='bases:sin_privilegios')
def producto_homologar(request, id):
    from catalogos.models import CatalogoSIN

    template_name = "inv/producto_homologar.html"
    prod = Producto.objects.filter(pk=id).first()
    if not prod:
        messages.error(request, 'Producto no existe.')
        return redirect('inv:producto_list')

    productos_sin = CatalogoSIN.objects.filter(
        tipo_catalogo=CatalogoSIN.TipoCatalogo.PRODUCTOS_SERVICIOS, vigente=True
    ).exclude(codigo_actividad__isnull=True).order_by('codigo_actividad', 'descripcion')

    if not productos_sin.exists():
        messages.error(
            request,
            'El catálogo de Productos/Servicios del SIN todavía no está sincronizado '
            '(o no trajo datos). Ejecute la sincronización de catálogos antes de homologar.'
        )
        return redirect('inv:producto_list')

    actividades_codigos = list(
        productos_sin.values_list('codigo_actividad', flat=True).distinct()
    )
    actividades = CatalogoSIN.objects.filter(
        tipo_catalogo=CatalogoSIN.TipoCatalogo.ACTIVIDADES,
        codigo__in=actividades_codigos,
        vigente=True,
    ).order_by('descripcion')
    codigos_con_descripcion = set(actividades.values_list('codigo', flat=True))
    actividades_faltantes = [c for c in actividades_codigos if c not in codigos_con_descripcion]

    if request.method == 'POST':
        actividad = request.POST.get('actividad_economica_sin')
        producto_sin_id = request.POST.get('codigo_producto_sin')

        if not actividad or not producto_sin_id:
            messages.error(request, 'Debe seleccionar actividad y producto/servicio SIN.')
            return redirect('inv:producto_homologar', id=id)

        item_sin = productos_sin.filter(codigo=producto_sin_id, codigo_actividad=actividad).first()
        if not item_sin:
            messages.error(request, 'La combinación seleccionada no es válida.')
            return redirect('inv:producto_homologar', id=id)

        prod.actividad_economica_sin = actividad
        prod.codigo_producto_sin = item_sin.codigo
        prod.save()

        messages.success(
            request,
            f'Producto "{prod.descripcion}" homologado: actividad {actividad}, '
            f'producto SIN {item_sin.codigo} - {item_sin.descripcion}.'
        )
        return redirect('inv:producto_list')

    # --- FIX: el template espera 'productos_sin_data' (un array de
    # dicts con claves 'actividad'/'codigo'/'descripcion') para armar
    # el select en cascada via JS + json_script. Antes solo se pasaba
    # 'productos_sin' (el queryset), asi que esa variable no existia
    # en el contexto del template -- Django la renderizaba como cadena
    # vacia, json_script serializaba "" en vez de un array, y el JS
    # rompia al intentar hacer .filter() sobre un string. Se arma aca
    # la lista con las claves correctas (el modelo usa
    # 'codigo_actividad', no 'actividad', que es lo que consume el JS).
    productos_sin_data = [
        {
            'actividad': p.codigo_actividad,
            'codigo': p.codigo,
            'descripcion': p.descripcion,
        }
        for p in productos_sin
    ]

    return render(request, template_name, {
        'producto': prod,
        'actividades': actividades,
        'actividades_faltantes': actividades_faltantes,
        'productos_sin': productos_sin,
        'productos_sin_data': productos_sin_data,
    })


@login_required(login_url='/login/')
@permission_required('inv.view_producto', login_url='bases:sin_privilegios')
def producto_homologar_pendientes(request):
    productos = Producto.objects.filter(estado=True).select_related('unidad_medida')
    pendientes = [p for p in productos if not p.homologado_sin]
    return render(request, 'inv/producto_homologar_pendientes.html', {
        'pendientes': pendientes,
        'total_productos': productos.count(),
        'total_pendientes': len(pendientes),
    })


# =====================================================================
# Tipo de Cambio y Revision de Precios (productos importados)
# =====================================================================

class TipoCambioView(SinPrivilegios, generic.ListView):
    permission_required = "inv.gestionar_precios_tc"
    model = TipoCambio
    template_name = "inv/tipo_cambio_list.html"
    context_object_name = "obj"

class TipoCambioNew(SuccessMessageMixin, SinPrivilegios, generic.CreateView):
    permission_required = "inv.gestionar_precios_tc"
    model = TipoCambio
    template_name = "inv/tipo_cambio_form.html"
    context_object_name = "obj"
    form_class = TipoCambioForm
    success_url = reverse_lazy("inv:tipo_cambio_list")
    success_message = "Tipo de Cambio Registrado Satisfactoriamente"

    def form_valid(self, form):
        form.instance.uc = self.request.user
        return super().form_valid(form)

class TipoCambioEdit(SuccessMessageMixin, SinPrivilegios, generic.UpdateView):
    permission_required = "inv.gestionar_precios_tc"
    model = TipoCambio
    template_name = "inv/tipo_cambio_form.html"
    context_object_name = "obj"
    form_class = TipoCambioForm
    success_url = reverse_lazy("inv:tipo_cambio_list")
    success_message = "Tipo de Cambio Actualizado Satisfactoriamente"

    def form_valid(self, form):
        form.instance.um = self.request.user.id
        return super().form_valid(form)


@login_required(login_url='/login/')
@permission_required('inv.gestionar_precios_sucursal', login_url='bases:sin_privilegios')
def precios_sucursal(request):
    """
    Precio de venta local por sucursal (25/09/2026, pedido de Carlos).
    Solo Administrador/Supervisor (inv.gestionar_precios_sucursal).
    Se edita el precio de UNA sucursal a la vez (las que el usuario puede
    ver, sin la Central: la Central usa el precio base). Campo vacio =
    sin precio local, la sucursal usa el precio base. Guardar valida todo
    antes de tocar nada, y deja rastro en el historial de precios.
    """
    from fe.models import Sucursal
    from bases.alcance import sucursales_visibles_ids, sucursal_elegida

    # Solo la Central fija precios (25/09/2026, pedido de Carlos): una
    # sucursal no cambia ningun precio, ni el suyo.
    if not es_casa_matriz(request):
        messages.error(request, _MSJ_PRECIOS_SOLO_CENTRAL)
        return redirect('inv:producto_list')

    sucursales = Sucursal.objects.exclude(codigo_sucursal=0).order_by('codigo_sucursal')
    ids = sucursales_visibles_ids(request.user)
    if ids is not None:
        sucursales = sucursales.filter(pk__in=ids)
    sucursales = list(sucursales)
    if not sucursales:
        return render(request, 'inv/precios_sucursal.html', {'sucursales': [], 'sin_sucursales': True})

    pedido = request.POST.get('sucursal_id') if request.method == 'POST' else None
    elegida_id = int(pedido) if pedido and pedido.isdigit() else sucursal_elegida(request)
    actual = obtener_sucursal_actual(request)
    sucursal = next((s for s in sucursales if s.pk == elegida_id), None) \
        or next((s for s in sucursales if actual and s.pk == actual.pk), None) \
        or sucursales[0]

    productos = list(Producto.objects.filter(estado=True).order_by('codigo'))
    locales = {
        r.producto_id: r for r in PrecioSucursal.objects.filter(sucursal=sucursal)
    }

    if request.method == 'POST':
        nuevos, errores = {}, []
        for p in productos:
            crudo = (request.POST.get(f'precio_{p.id}') or '').strip().replace(',', '.')
            if crudo == '':
                nuevos[p.id] = None
                continue
            try:
                valor = float(crudo)
                if valor <= 0:
                    raise ValueError
            except ValueError:
                errores.append(f'{p.codigo}: "{crudo}" no es un precio válido (debe ser mayor a 0).')
                continue
            nuevos[p.id] = round(valor, 2)
        if errores:
            for e in errores[:10]:
                messages.error(request, e)
        else:
            cambios = 0
            with transaction.atomic():
                for p in productos:
                    nuevo = nuevos.get(p.id)
                    previo = locales[p.id].precio if p.id in locales else None
                    if nuevo == previo:
                        continue
                    if nuevo is None:
                        locales[p.id].delete()
                    else:
                        PrecioSucursal.objects.update_or_create(
                            producto=p, sucursal=sucursal, defaults={'precio': nuevo})
                    HistorialPrecioProducto.objects.create(
                        producto=p, sucursal=sucursal,
                        precio_anterior=previo if previo is not None else p.precio,
                        precio_nuevo=nuevo if nuevo is not None else p.precio,
                        motivo=f'Precio local {sucursal.nombre}'
                               + (' (quitado, vuelve al precio base)' if nuevo is None else ''),
                    )
                    cambios += 1
            messages.success(request, f'Precios de {sucursal.nombre} guardados ({cambios} cambio(s)).')
            return redirect(f"{reverse_lazy('inv:precios_sucursal')}?sucursal={sucursal.pk}")

    filas = [{'producto': p, 'local': locales[p.id].precio if p.id in locales else None} for p in productos]
    return render(request, 'inv/precios_sucursal.html', {
        'sucursales': sucursales, 'sucursal': sucursal, 'filas': filas,
    })


def _precios_solo_lectura(request):
    """Revision de Precios se ANALIZA y EJECUTA solo desde la Central (Casa
    Matriz, codigo_sucursal=0) -- Producto.precio es unico para toda la
    empresa, asi que un cambio hecho desde una sucursal afectaria a todas.
    Desde cualquier otra sucursal la pantalla es de solo lectura
    (25/09/2026, pedido de Carlos). Sin sucursal resuelta (instalacion de
    una sola sucursal o sin sucursales) no se restringe nada."""
    return not es_casa_matriz(request)


_MSJ_PRECIOS_SOLO_CENTRAL = (
    'Los precios solo se pueden actualizar desde la Central (Casa Matriz). '
    'Desde una sucursal esta pantalla es de solo lectura.'
)


@login_required(login_url='/login/')
@permission_required('inv.gestionar_precios_tc', login_url='bases:sin_privilegios')
def revision_precios(request):
    UMBRAL_VARIACION_PCT = 3

    tipo_cambio_actual = TipoCambio.objects.filter(estado=True).order_by('-fecha').first()
    productos = Producto.objects.filter(
        estado=True, precio_referencia_usd__isnull=False
    ).order_by('codigo')

    filas = []
    for p in productos:
        precio_sugerido = p.calcular_precio_sugerido(tipo_cambio_actual)
        variacion = p.variacion_tipo_cambio_pct(tipo_cambio_actual)
        filas.append({
            'producto': p,
            'precio_sugerido': precio_sugerido,
            'variacion_pct': variacion,
            'supera_umbral': variacion is not None and abs(variacion) >= UMBRAL_VARIACION_PCT,
        })

    return render(request, 'inv/revision_precios.html', {
        'filas': filas,
        'tipo_cambio_actual': tipo_cambio_actual,
        'umbral': UMBRAL_VARIACION_PCT,
        'solo_lectura': _precios_solo_lectura(request),
    })


def _aplicar_precio_sugerido(producto, tipo_cambio_actual, usuario):
    precio_sugerido = producto.calcular_precio_sugerido(tipo_cambio_actual)
    if precio_sugerido is None:
        return None

    HistorialPrecioProducto.objects.create(
        producto=producto,
        precio_anterior=producto.precio,
        precio_nuevo=precio_sugerido,
        tipo_cambio_usado=tipo_cambio_actual,
        motivo="Ajuste por tipo de cambio",
        uc=usuario,
    )
    producto.precio = precio_sugerido
    producto.tipo_cambio_referencia = tipo_cambio_actual
    producto.save()
    return precio_sugerido


@login_required(login_url='/login/')
@permission_required('inv.gestionar_precios_tc', login_url='bases:sin_privilegios')
def aplicar_precio_sugerido(request, id):
    if request.method != 'POST':
        return redirect('inv:revision_precios')
    if _precios_solo_lectura(request):
        messages.error(request, _MSJ_PRECIOS_SOLO_CENTRAL)
        return redirect('inv:revision_precios')

    producto = Producto.objects.filter(pk=id).first()
    if not producto:
        messages.error(request, 'Producto no existe.')
        return redirect('inv:revision_precios')

    tipo_cambio_actual = TipoCambio.objects.filter(estado=True).order_by('-fecha').first()
    precio_nuevo = _aplicar_precio_sugerido(producto, tipo_cambio_actual, request.user)

    if precio_nuevo is None:
        messages.error(request, 'No se pudo calcular un precio sugerido para este producto.')
    else:
        messages.success(request, f'Precio de "{producto.descripcion}" actualizado a Bs {precio_nuevo}.')

    return redirect('inv:revision_precios')


@login_required(login_url='/login/')
@permission_required('inv.gestionar_precios_tc', login_url='bases:sin_privilegios')
def aplicar_todos_sugeridos(request):
    if request.method != 'POST':
        return redirect('inv:revision_precios')
    if _precios_solo_lectura(request):
        messages.error(request, _MSJ_PRECIOS_SOLO_CENTRAL)
        return redirect('inv:revision_precios')

    UMBRAL_VARIACION_PCT = 3
    tipo_cambio_actual = TipoCambio.objects.filter(estado=True).order_by('-fecha').first()
    productos = Producto.objects.filter(estado=True, precio_referencia_usd__isnull=False)

    aplicados = 0
    for p in productos:
        variacion = p.variacion_tipo_cambio_pct(tipo_cambio_actual)
        if variacion is not None and abs(variacion) >= UMBRAL_VARIACION_PCT:
            if _aplicar_precio_sugerido(p, tipo_cambio_actual, request.user) is not None:
                aplicados += 1

    if aplicados:
        messages.success(request, f'{aplicados} producto(s) actualizados por lote.')
    else:
        messages.warning(request, 'Ningún producto superaba el umbral de variación para aplicar en lote.')

    return redirect('inv:revision_precios')


# =====================================================================
# Ajuste de Inventario (13/09/2026) -- ver instructivo
# ajuste_inventario_instructivo.md. Dos flujos:
#  1. Produccion Interna: pantalla interactiva aca abajo (mismo patron
#     que compras() en cmp/views.py).
#  2. Carga Inicial: exportar/importar Excel, mas abajo del todo.
# =====================================================================

class AjusteInventarioView(SinPrivilegios, generic.ListView):
    permission_required = "inv.view_ajusteinventarioenc"
    model = AjusteInventarioEnc
    template_name = "inv/ajuste_inventario_list.html"
    context_object_name = "obj"

    def get_queryset(self):
        from bases.alcance import filtrar_por_sucursal
        return filtrar_por_sucursal(
            AjusteInventarioEnc.objects.filter(estado=True).select_related('motivo').order_by('-id'),
            self.request,
        )

    def get_context_data(self, **kwargs):
        from bases.alcance import contexto_filtro_sucursal
        context = super().get_context_data(**kwargs)
        context.update(contexto_filtro_sucursal(self.request))
        return context


@login_required(login_url='/login/')
@permission_required('inv.change_ajusteinventarioenc', login_url='bases:sin_privilegios')
@requiere_alcance('inv.AjusteInventarioEnc', ('sucursal_id',), 'ajuste_id')
def ajuste_inventario(request, ajuste_id=None):
    """
    Pantalla de Produccion Interna -- misma mecanica de interaccion
    que compras() (elegir producto, cargar cantidad y costo, agregar
    al detalle), pero con Motivo en vez de Proveedor y sin No.
    Factura/Fecha Factura. El motivo queda restringido a
    'es_produccion_interna=True' por AjusteInventarioEncForm -- la
    Carga Inicial NUNCA se crea desde esta pantalla manual.

    Mismo criterio que compras(): todas las validaciones ANTES de
    crear/actualizar el encabezado, para no dejar una cabecera
    huerfana si algo falla despues.
    """
    template_name = "inv/ajuste_inventario.html"
    productos = Producto.objects.filter(estado=True)

    if request.method == 'GET':
        enc = AjusteInventarioEnc.objects.filter(pk=ajuste_id).first()

        if enc:
            det = AjusteInventarioDet.objects.filter(ajuste=enc).order_by('id')
            form_enc = AjusteInventarioEncForm({
                'fecha': datetime.date.isoformat(enc.fecha),
                'motivo': enc.motivo_id,
                'observacion': enc.observacion,
            })
        else:
            det = None
            form_enc = AjusteInventarioEncForm(initial={'fecha': datetime.date.today()})

        if enc:
            numero_ajuste = enc.id
        else:
            ultimo = AjusteInventarioEnc.objects.order_by('-id').first()
            numero_ajuste = (ultimo.id + 1) if ultimo else 1

        contexto = {
            'productos': productos,
            'encabezado': enc,
            'detalle': det,
            'form_enc': form_enc,
            'numero_ajuste': numero_ajuste,
            'total_detalle': sum(d.sub_total for d in det) if det else 0,
        }
        return render(request, template_name, contexto)

    if request.method == 'POST':
        fecha = request.POST.get("fecha")
        motivo_id = request.POST.get("motivo")
        observacion = request.POST.get("observacion")

        motivo = MotivoAjusteInventario.objects.filter(
            pk=motivo_id, estado=True, es_produccion_interna=True
        ).first()
        if not motivo:
            messages.error(request, 'Debe seleccionar un motivo válido (Producción Interna).')
            return redirect("inv:ajuste_inventario_list")

        # Sucursal (Fase 2, 20/09/2026): se resuelve sola segun quien
        # esta logueado -- no se le pide elegir al cajero/almacenero.
        sucursal_actual = obtener_sucursal_actual(request)
        if sucursal_actual is None:
            messages.error(
                request,
                'No se pudo determinar su sucursal. Pida a un Administrador que se la '
                'asigne en Usuarios y Roles antes de registrar un ajuste de inventario.'
            )
            return redirect("inv:ajuste_inventario_list")

        # Produccion Interna restringida a Casa Matriz (Fase 2, diseño
        # confirmado por Carlos 16/09/2026) -- motivo.es_produccion_interna
        # ya esta garantizado True por el filtro de arriba (este modulo
        # SOLO permite motivos de Produccion Interna, la Carga Inicial
        # es aparte). codigo_sucursal=0 es la convencion ya establecida
        # en todo el sistema para "Casa Matriz" (ver fe/models.py).
        if sucursal_actual.codigo_sucursal != 0:
            messages.error(
                request,
                f'Producción Interna solo se puede registrar desde Casa Matriz '
                f'(su sucursal actual es "{sucursal_actual}").'
            )
            return redirect("inv:ajuste_inventario_list")

        producto_id = request.POST.get("id_id_producto")
        cantidad = request.POST.get("id_cantidad_detalle")
        costo_unitario = request.POST.get("id_costo_detalle")

        prod = Producto.objects.filter(pk=producto_id).first()
        if not prod:
            messages.error(request, 'El producto seleccionado no existe o no es válido')
            return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id) if ajuste_id else redirect("inv:ajuste_inventario_list")

        try:
            cantidad_num = int(float(cantidad))
            costo_num = float(costo_unitario)
        except (TypeError, ValueError):
            messages.error(request, 'Datos de cantidad/costo inválidos.')
            return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id) if ajuste_id else redirect("inv:ajuste_inventario_list")

        if cantidad_num <= 0:
            messages.error(request, 'La cantidad debe ser mayor a 0.')
            return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id) if ajuste_id else redirect("inv:ajuste_inventario_list")

        if costo_num <= 0:
            messages.error(request, 'El costo unitario debe ser mayor a 0.')
            return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id) if ajuste_id else redirect("inv:ajuste_inventario_list")

        if not ajuste_id:
            enc = AjusteInventarioEnc(
                fecha=fecha, motivo=motivo, observacion=observacion, uc=request.user,
                sucursal=sucursal_actual,
            )
            enc.save()
            ajuste_id = enc.id
        else:
            enc = AjusteInventarioEnc.objects.filter(pk=ajuste_id).first()
            if not enc:
                messages.error(request, 'El ajuste de inventario no existe.')
                return redirect("inv:ajuste_inventario_list")
            enc.fecha = fecha
            enc.motivo = motivo
            enc.observacion = observacion
            enc.um = request.user.id
            enc.save()

        AjusteInventarioDet.objects.create(
            ajuste=enc, producto=prod, cantidad=cantidad_num,
            costo_unitario=costo_num, uc=request.user
        )

        return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id)

    return render(request, template_name, {})


@login_required(login_url='/login/')
@permission_required('inv.delete_ajusteinventariodet', login_url='bases:sin_privilegios')
@requiere_alcance('inv.AjusteInventarioEnc', ('sucursal_id',), 'ajuste_id')
def ajuste_inventario_det_eliminar(request, ajuste_id, pk):
    """
    Quita una linea de un Ajuste de Inventario. A diferencia de
    Compras (que pasa a "contra compra" para no perder el rastro), acá
    se borra de verdad: esta pantalla es de uso reciente/frecuente
    (Produccion Interna) para corregir una carga mal hecha en el
    momento, no un documento fiscal ni contable que deba dejar huella
    de reversion. El post_delete (detalle_ajuste_borrar) revierte el
    stock solo.

    Vista de FUNCION a proposito, no DeleteView -- DeleteView.delete()
    es codigo muerto en Django 5.2 (ver CompraDetDelete en
    cmp/views.py para el hallazgo completo); una vista de funcion no
    tiene ese problema.
    """
    det = AjusteInventarioDet.objects.filter(pk=pk, ajuste_id=ajuste_id).first()
    if not det:
        messages.error(request, 'El detalle no existe.')
        return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id)

    if request.method == 'POST':
        prod = det.producto
        # Stock de la sucursal de ESTE ajuste (Fase 2, 20/09/2026), no
        # el total agregado de la empresa -- otra sucursal puede tener
        # de sobra mientras esta especifica quedaria en negativo.
        stock_en_sucursal = StockSucursal.objects.filter(
            producto=prod, sucursal=det.ajuste.sucursal
        ).first()
        cantidad_actual = stock_en_sucursal.cantidad if stock_en_sucursal else int(prod.existencia)
        if cantidad_actual - int(det.cantidad) < 0:
            messages.error(
                request,
                f'No se puede quitar: "{prod.descripcion}" quedaría con stock negativo en '
                f'{det.ajuste.sucursal or "esta sucursal"} (actual: {cantidad_actual}, se revertirían {det.cantidad}).'
            )
            return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id)

        det.delete()
        messages.success(request, 'Línea eliminada correctamente.')
        return redirect("inv:ajuste_inventario_edit", ajuste_id=ajuste_id)

    return render(request, "inv/ajuste_inventario_det_eliminar.html", {"obj": det})


@login_required(login_url='/login/')
@permission_required('inv.eliminar_ajusteinventarioenc', login_url='bases:sin_privilegios')
@requiere_alcance('inv.AjusteInventarioEnc', ('sucursal_id',), 'id')
def eliminar_ajuste_inventario(request, id):
    """Elimina (soft-delete) un Ajuste de Inventario completo, revirtiendo
    el stock de cada linea -- mismo patron que eliminar_compra."""
    enc = AjusteInventarioEnc.objects.filter(pk=id).first()
    if not enc:
        messages.error(request, 'El ajuste de inventario no existe.')
        return redirect('inv:ajuste_inventario_list')

    if not enc.estado:
        messages.error(request, 'Este ajuste ya fue eliminado.')
        return redirect('inv:ajuste_inventario_list')

    if request.method == 'POST':
        detalles = AjusteInventarioDet.objects.filter(ajuste=enc)

        cantidad_por_producto = {}
        for det in detalles:
            cantidad_por_producto[det.producto_id] = cantidad_por_producto.get(det.producto_id, 0) + det.cantidad

        # Stock de la sucursal de ESTE ajuste (Fase 2, 20/09/2026), no
        # el total agregado de la empresa.
        productos_en_negativo = []
        for producto_id, cantidad_total in cantidad_por_producto.items():
            prod = Producto.objects.get(pk=producto_id)
            stock_en_sucursal = StockSucursal.objects.filter(
                producto=prod, sucursal=enc.sucursal
            ).first()
            cantidad_actual = stock_en_sucursal.cantidad if stock_en_sucursal else int(prod.existencia)
            if cantidad_actual - cantidad_total < 0:
                productos_en_negativo.append(
                    f'{prod.descripcion} (stock actual en {enc.sucursal or "esta sucursal"}: '
                    f'{cantidad_actual}, se revertirían {cantidad_total})'
                )

        if productos_en_negativo:
            messages.error(
                request,
                'No se puede eliminar: dejaría stock negativo en: ' + '; '.join(productos_en_negativo) +
                '. Probablemente ya se vendió parte de esta mercadería -- revise antes de continuar.'
            )
            return redirect('inv:ajuste_inventario_edit', ajuste_id=id)

        enc.estado = False
        enc.save()

        for producto_id, cantidad_total in cantidad_por_producto.items():
            ajustar_stock_sucursal(producto_id, enc.sucursal, -cantidad_total)

        messages.success(request, 'Ajuste de inventario eliminado correctamente.')
        return redirect('inv:ajuste_inventario_list')

    return render(request, 'inv/ajuste_inventario_eliminar.html', {'enc': enc})


# =====================================================================
# Carga Inicial por Excel (13/09/2026) -- ver seccion 5 del
# instructivo. Exportar plantilla -> completar -> importar con
# validacion previa completa (todo o nada).
# =====================================================================

CARGA_INICIAL_COLUMNAS = [
    "Código", "Descripción", "Marca", "Categoría", "Subcategoría", "Unidad de Medida",
    "Cantidad", "Costo Unitario", "Precio de Venta",
    "Precio Referencial USD", "Margen Deseado %",
]


@login_required(login_url='/login/')
@permission_required('inv.cargar_inventario_inicial', login_url='bases:sin_privilegios')
def carga_inicial(request):
    """Pantalla de Carga Inicial de Inventario: exportar plantilla / importar Excel."""
    return render(request, 'inv/carga_inicial.html', {})


@login_required(login_url='/login/')
@permission_required('inv.cargar_inventario_inicial', login_url='bases:sin_privilegios')
def carga_inicial_exportar_plantilla(request):
    """
    Genera la plantilla: una fila por cada producto YA existente
    (Código y Descripción rellenos, el resto vacio para completar),
    mas filas vacias al final por si se quieren agregar productos
    nuevos. NO crea Marca/Categoría/Subcategoria/Unidad de Medida --
    solo las busca por nombre al importar (mensaje de aviso en la
    plantilla HTML, ver carga_inicial.html).

    Columna Categoría (agregada 14/09/2026): se pide ADEMAS de
    Subcategoría, no en su lugar -- 'SubCategoria.descripcion' no es
    unica por si sola (dos categorias distintas podrian tener una
    subcategoria con el mismo nombre), asi que buscar la subcategoria
    solo por nombre seria ambiguo. Con la Categoria como referencia,
    la busqueda al importar es (categoria, subcategoria) -- ver
    carga_inicial_importar.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Carga Inicial"

    ws.append(CARGA_INICIAL_COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for p in Producto.objects.filter(estado=True).order_by('codigo'):
        ws.append([p.codigo, p.descripcion, '', '', '', '', '', '', '', '', ''])
        ws.cell(row=ws.max_row, column=1).number_format = '@'

    for _ in range(20):
        ws.append([])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="plantilla_carga_inicial_inventario.xlsx"'
    wb.save(response)
    return response


@login_required(login_url='/login/')
@permission_required('inv.cargar_inventario_inicial', login_url='bases:sin_privilegios')
def carga_inicial_importar(request):
    """
    Importa la plantilla completa. Valida fila por fila, TODO antes de
    tocar nada (mismo criterio de "no dejar nada a medio camino" que
    ya rige en todo el sistema) -- si CUALQUIER fila falla, no se
    importa NADA y se muestra la lista completa de errores.

    NOTA (desviacion del instructivo, a confirmar con Carlos): la
    plantilla no trae una columna de Código de Barra, pero
    Producto.codigo_barra es obligatorio a nivel de modelo. Para un
    producto NUEVO creado por esta via, se usa el mismo 'Código' como
    codigo_barra provisorio -- se puede corregir despues a mano desde
    el formulario de Producto si se consigue el codigo de barra real.
    """
    if request.method != 'POST':
        return redirect('inv:carga_inicial')

    archivo = request.FILES.get('archivo')
    if not archivo:
        messages.error(request, 'Debe seleccionar un archivo Excel (.xlsx).')
        return redirect('inv:carga_inicial')

    # Sucursal (25/09/2026): la carga inicial suma stock a la sucursal
    # ACTUAL del usuario -- antes el ajuste se guardaba sin sucursal y el
    # stock solo entraba al total de la empresa (no a StockSucursal), asi
    # que esa sucursal no podia vender lo cargado ni transferirlo.
    sucursal_actual = obtener_sucursal_actual(request)
    if sucursal_actual is None:
        messages.error(
            request,
            'No se pudo determinar su sucursal. Pida a un Administrador que se la '
            'asigne en Usuarios y Roles antes de hacer la carga inicial.'
        )
        return redirect('inv:carga_inicial')

    try:
        wb = load_workbook(archivo, data_only=True)
        ws = wb.active
    except Exception:
        messages.error(request, 'No se pudo leer el archivo. Verifique que sea un Excel (.xlsx) válido.')
        return redirect('inv:carga_inicial')

    marcas_por_nombre = {m.descripcion.upper(): m for m in Marca.objects.filter(estado=True)}
    categorias_por_nombre = {c.descripcion.upper(): c for c in Categoria.objects.filter(estado=True)}
    # Clave (categoria_id, nombre) -- 'SubCategoria.descripcion' NO es
    # unica por si sola (dos categorias distintas pueden tener una
    # subcategoria con el mismo nombre), asi que buscarla solo por
    # nombre seria ambiguo. Se busca por la PAREJA categoria+nombre.
    subcats_por_categoria_y_nombre = {
        (s.categoria_id, s.descripcion.upper()): s for s in SubCategoria.objects.filter(estado=True)
    }
    ums_por_nombre = {u.descripcion.upper(): u for u in UnidadMedida.objects.filter(estado=True)}
    productos_por_codigo = {p.codigo.upper(): p for p in Producto.objects.all()}

    errores = []
    filas_validas = []
    fila_num = 1

    for row in ws.iter_rows(min_row=2, max_col=11, values_only=True):
        fila_num += 1
        if all(v in (None, '') for v in row):
            continue

        (codigo, descripcion, marca_nom, categoria_nom, subcat_nom, um_nom,
         cantidad, costo, precio, precio_ref, margen) = (list(row) + [None] * 11)[:11]

        if not codigo:
            errores.append(f"fila {fila_num}: falta el Código.")
            continue
        codigo = str(codigo).strip().upper()
        producto_existente = productos_por_codigo.get(codigo)

        cantidad_num = None
        try:
            cantidad_num = int(float(cantidad))
            if cantidad_num <= 0:
                errores.append(f"fila {fila_num}: la Cantidad debe ser mayor a 0.")
        except (TypeError, ValueError):
            errores.append(f"fila {fila_num}: la Cantidad no es un número válido.")

        costo_num = None
        try:
            costo_num = float(costo)
            if costo_num <= 0:
                errores.append(f"fila {fila_num}: el Costo Unitario debe ser mayor a 0.")
        except (TypeError, ValueError):
            errores.append(f"fila {fila_num}: el Costo Unitario no es un número válido.")

        marca_obj = categoria_obj = subcat_obj = um_obj = None
        if not producto_existente:
            if not es_casa_matriz(request):
                errores.append(
                    f"fila {fila_num}: el código '{codigo}' no existe -- los productos nuevos "
                    "solo se pueden crear desde la Central."
                )
                continue
            if not descripcion:
                errores.append(f"fila {fila_num}: falta la Descripción (producto nuevo, código '{codigo}').")
            if not precio:
                errores.append(f"fila {fila_num}: falta el Precio de Venta (obligatorio para un producto nuevo, código '{codigo}').")

            if not marca_nom:
                errores.append(f"fila {fila_num}: falta la Marca (producto nuevo, código '{codigo}').")
            else:
                marca_obj = marcas_por_nombre.get(str(marca_nom).strip().upper())
                if not marca_obj:
                    errores.append(f"fila {fila_num}: la marca '{marca_nom}' no existe, debe crearla primero en Catálogos.")

            # Categoría se valida ANTES que Subcategoría -- la
            # subcategoria se busca por (categoria, nombre), asi que
            # sin una categoria valida no tiene sentido buscarla.
            if not categoria_nom:
                errores.append(f"fila {fila_num}: falta la Categoría (producto nuevo, código '{codigo}').")
            else:
                categoria_obj = categorias_por_nombre.get(str(categoria_nom).strip().upper())
                if not categoria_obj:
                    errores.append(f"fila {fila_num}: la categoría '{categoria_nom}' no existe, debe crearla primero en Catálogos.")

            if not subcat_nom:
                errores.append(f"fila {fila_num}: falta la Subcategoría (producto nuevo, código '{codigo}').")
            elif categoria_obj:
                subcat_obj = subcats_por_categoria_y_nombre.get((categoria_obj.id, str(subcat_nom).strip().upper()))
                if not subcat_obj:
                    errores.append(
                        f"fila {fila_num}: la subcategoría '{subcat_nom}' no existe bajo la categoría "
                        f"'{categoria_nom}', debe crearla primero en Catálogos."
                    )
            # Si la categoria no existia, ya se reporto ese error arriba
            # -- no hace falta duplicar el reclamo sobre la subcategoria
            # (no se puede validar sin saber antes bajo que categoria).

            if not um_nom:
                errores.append(f"fila {fila_num}: falta la Unidad de Medida (producto nuevo, código '{codigo}').")
            else:
                um_obj = ums_por_nombre.get(str(um_nom).strip().upper())
                if not um_obj:
                    errores.append(f"fila {fila_num}: la unidad de medida '{um_nom}' no existe, debe crearla primero en Catálogos.")

        filas_validas.append(dict(
            fila=fila_num, codigo=codigo, descripcion=descripcion,
            producto_existente=producto_existente, marca=marca_obj,
            subcategoria=subcat_obj, unidad_medida=um_obj,
            cantidad=cantidad_num, costo=costo_num,
            precio=precio, precio_ref=precio_ref, margen=margen,
        ))

    if errores:
        return render(request, 'inv/carga_inicial.html', {'errores': errores})

    if not filas_validas:
        messages.warning(request, 'El archivo no tiene ninguna fila para importar.')
        return redirect('inv:carga_inicial')

    motivo = MotivoAjusteInventario.objects.filter(descripcion='CARGA INICIAL').first()
    if not motivo:
        messages.error(
            request,
            'No existe el motivo "Carga Inicial" en el catálogo de Motivos de Ajuste de '
            'Inventario. Contacte al administrador.'
        )
        return redirect('inv:carga_inicial')

    with transaction.atomic():
        enc = AjusteInventarioEnc.objects.create(
            fecha=timezone.localdate(), motivo=motivo,
            observacion='Carga inicial de inventario por Excel', uc=request.user,
            sucursal=sucursal_actual,
        )
        creados = 0
        for f in filas_validas:
            prod = f['producto_existente']
            if not prod:
                prod = Producto(
                    codigo=f['codigo'],
                    codigo_barra=f['codigo'],
                    descripcion=f['descripcion'],
                    marca=f['marca'],
                    subcategoria=f['subcategoria'],
                    unidad_medida=f['unidad_medida'],
                    precio=float(f['precio']),
                    precio_referencia_usd=float(f['precio_ref']) if f['precio_ref'] else None,
                    margen_deseado_pct=float(f['margen']) if f['margen'] else 0,
                    uc=request.user,
                )
                prod.save()
                creados += 1
            else:
                cambios = False
                if f['precio']:
                    prod.precio = float(f['precio'])
                    cambios = True
                if f['precio_ref']:
                    prod.precio_referencia_usd = float(f['precio_ref'])
                    cambios = True
                if f['margen']:
                    prod.margen_deseado_pct = float(f['margen'])
                    cambios = True
                if cambios:
                    prod.save()

            AjusteInventarioDet.objects.create(
                ajuste=enc, producto=prod, cantidad=f['cantidad'],
                costo_unitario=f['costo'], uc=request.user
            )

    messages.success(
        request,
        f'Carga inicial importada correctamente: {len(filas_validas)} línea(s) cargadas, '
        f'{creados} producto(s) nuevo(s) creados (Ajuste #{enc.id}).'
    )
    # Redirige al listado, no a la pantalla de edicion (esa pantalla es
    # especifica de Produccion Interna -- su combo de Motivo solo
    # ofrece motivos con es_produccion_interna=True, y Carga Inicial no
    # lo es).
    return redirect('inv:ajuste_inventario_list')


# =====================================================================
# Transferencias de Stock entre Sucursales (Fase 2, 20/09/2026)
# =====================================================================

class TransferenciaStockListView(SinPrivilegios, generic.ListView):
    """
    Lista las transferencias donde la sucursal actual participa (como
    origen O como destino) -- un superusuario ve todas. Mismo criterio
    que el resto del sistema: nadie deberia ver movimientos de una
    sucursal que no es la suya, salvo Administrador.
    """
    permission_required = "inv.view_transferenciastockenc"
    model = TransferenciaStockEnc
    template_name = "inv/transferencia_stock_list.html"
    context_object_name = "obj"

    def get_queryset(self):
        from django.db.models import Q
        # Ojo: en este modelo 'estado' esta sobreescrito como el estado
        # del negocio (en_transito/confirmada/cancelada), no como el
        # booleano de soft-delete de ClaseModelo2 -- no filtrar por
        # estado=True aca, nunca matchearia nada.
        qs = TransferenciaStockEnc.objects.select_related(
            'sucursal_origen', 'sucursal_destino'
        ).order_by('-id')
        # Alcance por sucursal (24/09/2026): antes cualquier no-superusuario
        # veia solo las de su sucursal actual; ahora lo decide
        # PerfilUsuario.alcance (Todas / Solo su sucursal).
        from bases.alcance import q_alcance, sucursal_elegida
        qs = qs.filter(q_alcance(self.request.user, 'sucursal_origen', 'sucursal_destino'))
        elegida = sucursal_elegida(self.request)
        if elegida:
            qs = qs.filter(Q(sucursal_origen=elegida) | Q(sucursal_destino=elegida))
        return qs

    def get_context_data(self, **kwargs):
        from bases.alcance import contexto_filtro_sucursal
        context = super().get_context_data(**kwargs)
        context.update(contexto_filtro_sucursal(self.request))
        return context


@login_required(login_url='/login/')
@permission_required('inv.add_transferenciastockenc', login_url='bases:sin_privilegios')
@requiere_alcance('inv.TransferenciaStockEnc', ('sucursal_origen_id','sucursal_destino_id'), 'transferencia_id')
def transferencia_stock_new(request, transferencia_id=None):
    """
    Envia stock de la sucursal actual (origen, fija -- no se elige) a
    otra sucursal (destino, elegida en el formulario). Mismo patron de
    interaccion que compras()/ajuste_inventario(): elegir producto,
    cargar cantidad, agregar al detalle -- cada linea agregada
    descuenta stock del origen DE INMEDIATO (ver
    detalle_transferencia_guardar en inv/models.py); recien entra al
    destino cuando ese lado confirma la recepcion (transferencia_
    confirmar_recepcion, mas abajo).
    """
    from fe.models import Sucursal

    template_name = "inv/transferencia_stock_form.html"
    productos = Producto.objects.filter(estado=True)

    sucursal_actual = obtener_sucursal_actual(request)
    if sucursal_actual is None:
        messages.error(
            request,
            'No se pudo determinar su sucursal. Pida a un Administrador que se la '
            'asigne en Usuarios y Roles antes de registrar una transferencia.'
        )
        return redirect('inv:transferencia_stock_list')

    if request.method == 'GET':
        enc = TransferenciaStockEnc.objects.filter(pk=transferencia_id).first()
        det = TransferenciaStockDet.objects.filter(transferencia=enc).order_by('id') if enc else None

        if enc and enc.estado != TransferenciaStockEnc.EN_TRANSITO:
            # Una vez Confirmada o Cancelada, no se pueden agregar mas
            # lineas (ya se cerro el movimiento).
            messages.error(request, 'Esta transferencia ya no está en tránsito, no se pueden agregar líneas.')
            return redirect('inv:transferencia_stock_list')

        contexto = {
            'productos': productos,
            'encabezado': enc,
            'detalle': det,
            'sucursal_actual': sucursal_actual,
            'sucursales_destino': Sucursal.objects.exclude(pk=sucursal_actual.pk).order_by('codigo_sucursal'),
            'numero_transferencia': enc.id if enc else (
                (TransferenciaStockEnc.objects.order_by('-id').first().id + 1)
                if TransferenciaStockEnc.objects.exists() else 1
            ),
        }
        return render(request, template_name, contexto)

    if request.method == 'POST':
        producto_id = request.POST.get('id_id_producto')
        cantidad = request.POST.get('id_cantidad_detalle')

        prod = Producto.objects.filter(pk=producto_id).first()
        if not prod:
            messages.error(request, 'El producto seleccionado no existe o no es válido')
            return redirect('inv:transferencia_stock_edit', transferencia_id=transferencia_id) if transferencia_id else redirect('inv:transferencia_stock_new')

        try:
            cantidad_num = int(float(cantidad))
        except (TypeError, ValueError):
            messages.error(request, 'Cantidad inválida.')
            return redirect('inv:transferencia_stock_edit', transferencia_id=transferencia_id) if transferencia_id else redirect('inv:transferencia_stock_new')

        if cantidad_num <= 0:
            messages.error(request, 'La cantidad debe ser mayor a 0.')
            return redirect('inv:transferencia_stock_edit', transferencia_id=transferencia_id) if transferencia_id else redirect('inv:transferencia_stock_new')

        if not transferencia_id:
            sucursal_destino_id = request.POST.get('sucursal_destino')
            sucursal_destino = Sucursal.objects.filter(pk=sucursal_destino_id).first()
            if not sucursal_destino:
                messages.error(request, 'Debe elegir una sucursal destino válida.')
                return redirect('inv:transferencia_stock_new')
            if sucursal_destino.pk == sucursal_actual.pk:
                messages.error(request, 'La sucursal destino debe ser distinta de la actual.')
                return redirect('inv:transferencia_stock_new')
            enc = TransferenciaStockEnc(
                sucursal_origen=sucursal_actual, sucursal_destino=sucursal_destino, uc=request.user
            )
            enc.save()
            transferencia_id = enc.id
        else:
            enc = TransferenciaStockEnc.objects.filter(pk=transferencia_id).first()
            if not enc:
                messages.error(request, 'La transferencia no existe.')
                return redirect('inv:transferencia_stock_list')
            if enc.estado != TransferenciaStockEnc.EN_TRANSITO:
                messages.error(request, 'Esta transferencia ya no está en tránsito.')
                return redirect('inv:transferencia_stock_list')

        # Stock disponible en la sucursal de ORIGEN (Fase 2) -- no se
        # puede enviar mas de lo que hay fisicamente ahi.
        stock_en_origen = StockSucursal.objects.filter(producto=prod, sucursal=enc.sucursal_origen).first()
        disponible = stock_en_origen.cantidad if stock_en_origen else 0
        if cantidad_num > disponible:
            messages.error(
                request,
                f'No hay suficiente stock de "{prod.descripcion}" en {enc.sucursal_origen} '
                f'(disponible: {disponible}).'
            )
            return redirect('inv:transferencia_stock_edit', transferencia_id=transferencia_id)

        TransferenciaStockDet.objects.create(
            transferencia=enc, producto=prod, cantidad=cantidad_num, uc=request.user
        )

        return redirect('inv:transferencia_stock_edit', transferencia_id=transferencia_id)

    return render(request, template_name, {})


@login_required(login_url='/login/')
@permission_required('inv.confirmar_transferenciastock', login_url='bases:sin_privilegios')
@requiere_alcance('inv.TransferenciaStockEnc', ('sucursal_origen_id','sucursal_destino_id'), 'id')
def transferencia_confirmar_recepcion(request, id):
    """
    Confirma que la sucursal DESTINO recibio la mercaderia -- recien
    aca el stock entra de verdad a StockSucursal del destino (hasta
    este momento estaba "en transito", fuera de cualquier sucursal).
    Cualquiera con el permiso puede confirmar (no se restringe a que
    su sucursal_actual sea exactamente el destino -- un Administrador/
    Supervisor puede estar confirmando por telefono con el almacenero
    del destino, por ejemplo) -- la decision de quien tiene el permiso
    ya la toma crear_grupos_permisos.py.
    """
    enc = TransferenciaStockEnc.objects.filter(pk=id).first()
    if not enc:
        messages.error(request, 'La transferencia no existe.')
        return redirect('inv:transferencia_stock_list')

    if enc.estado != TransferenciaStockEnc.EN_TRANSITO:
        messages.error(request, f'Esta transferencia ya está "{enc.get_estado_display()}", no se puede confirmar de nuevo.')
        return redirect('inv:transferencia_stock_list')

    if request.method == 'POST':
        detalles = TransferenciaStockDet.objects.filter(transferencia=enc)
        if not detalles.exists():
            messages.error(request, 'Esta transferencia no tiene ningún producto cargado -- no hay nada que confirmar.')
            return redirect('inv:transferencia_stock_list')

        for det in detalles:
            ajustar_stock_sucursal(det.producto_id, enc.sucursal_destino, det.cantidad)

        enc.estado = TransferenciaStockEnc.CONFIRMADA
        enc.fecha_confirmacion = timezone.now()
        enc.usuario_confirmacion = request.user
        enc.save()

        messages.success(request, f'Transferencia #{enc.id} confirmada: stock sumado a {enc.sucursal_destino}.')
        return redirect('inv:transferencia_stock_list')

    return render(request, 'inv/transferencia_stock_confirmar.html', {'obj': enc})


@login_required(login_url='/login/')
@permission_required('inv.cancelar_transferenciastock', login_url='bases:sin_privilegios')
@requiere_alcance('inv.TransferenciaStockEnc', ('sucursal_origen_id','sucursal_destino_id'), 'id')
def transferencia_cancelar(request, id):
    """
    Cancela una transferencia todavia EN TRANSITO -- el stock que ya
    habia salido del origen (al agregar cada linea) vuelve a sumarse
    ahi. No se puede cancelar una ya Confirmada (para eso hace falta
    una transferencia nueva en sentido contrario, mismo criterio que
    "no se edita/borra un documento fiscal ya cerrado" del resto del
    sistema).
    """
    enc = TransferenciaStockEnc.objects.filter(pk=id).first()
    if not enc:
        messages.error(request, 'La transferencia no existe.')
        return redirect('inv:transferencia_stock_list')

    if enc.estado != TransferenciaStockEnc.EN_TRANSITO:
        messages.error(request, f'Esta transferencia ya está "{enc.get_estado_display()}", no se puede cancelar.')
        return redirect('inv:transferencia_stock_list')

    if request.method == 'POST':
        detalles = TransferenciaStockDet.objects.filter(transferencia=enc)
        for det in detalles:
            ajustar_stock_sucursal(det.producto_id, enc.sucursal_origen, det.cantidad)

        enc.estado = TransferenciaStockEnc.CANCELADA
        enc.save()

        messages.success(request, f'Transferencia #{enc.id} cancelada: stock devuelto a {enc.sucursal_origen}.')
        return redirect('inv:transferencia_stock_list')

    return render(request, 'inv/transferencia_stock_cancelar.html', {'obj': enc})